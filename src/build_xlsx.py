import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

D = json.load(open('/home/claude/portfolio/data.json'))
EPOCH = datetime.date(1970, 1, 1)
day2date = lambda d: EPOCH + datetime.timedelta(days=d)

F = 'Arial'
H_FONT = Font(name=F, bold=True, color='FFFFFF')
H_FILL = PatternFill('solid', fgColor='1F2A3C')
BLUE = Font(name=F, color='0000FF')
BLACK = Font(name=F)
GREEN = Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
NOTE = Font(name=F, italic=True, color='555555')
YELLOW = PatternFill('solid', fgColor='FFF2B3')
SEGFILL = {'Champions': 'DCE9FA', 'Potential Loyalists': 'DDF3E9', 'At Risk': 'FCE3D8', 'Hibernating': 'ECECEA'}
DATE = 'DD.MM.YYYY'
MONTH = 'MMM YYYY'
MONEY = '#,##0" ₸"'
PCT = '0.0%'

wb = Workbook()

def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = H_FONT; cell.fill = H_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

# ---------------- Шаги ----------------
ws = wb.active; ws.title = 'Шаги'
ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 30; ws.column_dimensions['C'].width = 110
ws['A1'] = 'RFM-анализ SmartFlow: откуда цифры и что я делала шаг за шагом'; ws['A1'].font = TITLE
ws['A2'] = 'Данные демонстрационные: сгенерированы моделью, которая имитирует покупки на SmartFlow. Реальные данные будут браться так же — из таблицы оплат в Supabase.'; ws['A2'].font = NOTE
steps = [
    ('0', 'Откуда данные', 'Модель имитирует 24 месяца покупок (12 месяцев «разгона» + 12 месяцев анализа). Каждый месяц приходит ≈3 300 новых покупателей. У каждого есть ниша (4 варианта), канал привлечения (4 варианта) и скрытый тип: лояльный (16%), средний (34%) или разовый (50%) — тип задаёт вероятность следующей покупки. Средний чек зависит от ниши (12–29 тыс. ₸). Интервал до следующей покупки — случайный, в среднем 20–33 дня в зависимости от ниши. 20 июля 2026 в модели меняется шаблон воронки: после первой оплаты Smart консультант больше не пишет — поэтому вторая покупка случается реже. Это и есть «падение повторных покупок за последние 2 месяца», которое мы ищем.'),
    ('1', 'Собрать заказы → лист «Заказы»', 'Беру все оплаты за 12 месяцев: с 01.10.2025 по 30.09.2026. Одна строка = одна оплата: номер заказа, покупатель, дата, сумма. Столбец «Новый» = 1, если это первая покупка человека вообще (у некоторых первая покупка была до октября 2025). Остальные столбцы — формулы: месяц заказа, первый ли это заказ покупателя в этом месяце, когорта покупателя.'),
    ('2', 'Свести к покупателю → лист «Покупатели»', 'Для каждого покупателя считаю три числа. F — сколько у него заказов за 12 месяцев. R — сколько дней прошло от последней покупки до даты анализа (30.09.2026). M — сумма всех его покупок. Заказы одного покупателя идут в «Заказах» подряд, поэтому формулы берут диапазон строк «с … по …».'),
    ('3', 'Подобрать пороги → лист «Пороги»', 'R: смотрю, через сколько дней люди делают вторую покупку. 90% повторных покупок случаются в пределах ≈56 дней → округляю вверх до 60 дней: кто молчит дольше, скорее всего ушёл. F: смотрю, какая доля людей с N покупками делает следующую. После 3-й покупки доля перестаёт расти и держится около 50% → «часто» = от 3 покупок. M: верхние 25% покупателей по сумме — «высокий M».'),
    ('4', 'Разделить на сегменты → лист «Сегменты»', 'По R и F получаю 4 сегмента: Champions (недавно и часто), Potential Loyalists (недавно, 1–2 покупки), At Risk (давно, но раньше часто), Hibernating (давно, 1–2 покупки). Для каждого считаю число людей, долю оборота, средние R, F, M и сколько в нём людей с высоким M.'),
    ('5', 'Динамика → лист «По месяцам»', 'По каждому месяцу: заказы, выручка, новые покупатели, повторные заказы, средний чек. Видно: новые покупатели приходят стабильно (≈3 100–3 500), а доля повторных заказов и «повтор за 30 дней» падают с июля–августа 2026.'),
    ('6', 'Когортный анализ → лист «Когорты»', 'Когорта = все, кто впервые купил в одном месяце. Для каждой когорты считаю, какая доля людей снова купила через 1, 2, 3… месяца. Если падение видно только у когорт с июля — проблема в том, что изменилось для новых покупателей, а не в рынке или сезоне.'),
    ('7', 'Пример вручную → лист «Пример»', '10 покупателей и их заказы. Те же R, F, M и сегмент посчитаны простыми формулами — чтобы можно было проверить логику руками.'),
    ('8', 'Стратегия → лист «Стратегия»', 'Перевожу выводы в деньги. Для каждого сегмента: сколько людей покупает в месяц сейчас и при цели → дополнительные покупатели × выручка на покупателя. Бюджет: раньше удержание делилось по числу людей (68% — на Hibernating), теперь — туда, где есть прирост. Итог: −20% бюджета и +11,7% выручки.'),
    ('9', 'Выводы', '1) Проблема после первой покупки: новых столько же, а повтор за 30 дней упал с ≈23% до ≈15%. 2) Падение во всех каналах и нишах → дело не в составе аудитории. 3) Падение начинается с покупок после 20 июля → причина — изменение воронки. 4) Отсюда стратегия: вернуть сообщения после первой оплаты (3-й, 7-й, 14-й день) для Potential Loyalists и персональные предложения для At Risk.'),
]
ws['A4'] = 'Шаг'; ws['B4'] = 'Что делаю'; ws['C4'] = 'Как и зачем'
for c in ('A4', 'B4', 'C4'):
    ws[c].font = H_FONT; ws[c].fill = H_FILL
for i, (n, t, d) in enumerate(steps, 5):
    ws.cell(row=i, column=1, value=n).font = BOLD
    ws.cell(row=i, column=2, value=t).font = BOLD
    c = ws.cell(row=i, column=3, value=d); c.font = BLACK; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(row=i, column=1).alignment = Alignment(vertical='top'); ws.cell(row=i, column=2).alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[i].height = max(30, 15 * (len(d) // 95 + 1))
r = len(steps) + 6
ws.cell(row=r, column=1, value='Цвета: синий текст — исходные данные и решения, которые можно менять; чёрный — формулы; жёлтая заливка — пороги.').font = NOTE

# ---------------- Заказы ----------------
wo = wb.create_sheet('Заказы')
header(wo, 1, ['Заказ', 'Покупатель', 'Дата', 'Сумма', 'Новый (1 = первая покупка вообще)', 'Месяц', 'Первый заказ покупателя в этом месяце', 'Когорта покупателя'],
       [11, 12, 12, 12, 16, 11, 18, 16])
first_ever = {c[0]: c[3] for c in D['customers']}
rows_of = {}
for i, (oid, cid, d, amt) in enumerate(D['orders'], 2):
    rows_of.setdefault(cid, [i, i])[1] = i
    is_new = 1 if d == first_ever[cid] and rows_of[cid][0] == i else 0
    wo.cell(row=i, column=1, value=oid).font = BLUE
    wo.cell(row=i, column=2, value=cid).font = BLUE
    c = wo.cell(row=i, column=3, value=day2date(d)); c.font = BLUE; c.number_format = DATE
    c = wo.cell(row=i, column=4, value=amt); c.font = BLUE; c.number_format = MONEY
    wo.cell(row=i, column=5, value=is_new).font = BLUE
    c = wo.cell(row=i, column=6, value=f'=DATE(YEAR(C{i}),MONTH(C{i}),1)'); c.number_format = MONTH
    wo.cell(row=i, column=7, value=f'=IF(OR(B{i}<>B{i-1},F{i}<>F{i-1}),1,0)')
    c = wo.cell(row=i, column=8, value=f'=IF(B{i}<>B{i-1},IF(E{i}=1,F{i},"до окт 2025"),H{i-1})'); c.number_format = MONTH
wo.freeze_panes = 'A2'
NO = len(D['orders']) + 1

# ---------------- Пороги (cells referenced by Покупатели) ----------------
wp = wb.create_sheet('Пороги')
wp.column_dimensions['A'].width = 52; wp.column_dimensions['B'].width = 16; wp.column_dimensions['C'].width = 16; wp.column_dimensions['D'].width = 16; wp.column_dimensions['E'].width = 60
wp['A1'] = 'Шаг 3. Пороги R, F и M'; wp['A1'].font = TITLE
wp['A3'] = 'Дата анализа'; wp['B3'] = datetime.date(2026, 9, 30); wp['B3'].number_format = DATE; wp['B3'].font = BLUE
wp['A4'] = 'Начало окна (12 месяцев)'; wp['B4'] = datetime.date(2025, 10, 1); wp['B4'].number_format = DATE; wp['B4'].font = BLUE

wp['A6'] = 'R — давность'; wp['A6'].font = BOLD
wp['A7'] = 'Покупателей со 2-й покупкой'; wp['B7'] = "=COUNT('Покупатели'!K:K)"
wp['A8'] = 'Медиана интервала 1-я → 2-я покупка, дней'; wp['B8'] = "=MEDIAN('Покупатели'!K:K)"
wp['A9'] = '75-й перцентиль, дней'; wp['B9'] = "=PERCENTILE('Покупатели'!K:K,0.75)"
wp['A10'] = '90-й перцентиль, дней'; wp['B10'] = "=PERCENTILE('Покупатели'!K:K,0.9)"
wp['A11'] = 'Порог R: 90-й перцентиль, вверх до 10 дней'; wp['B11'] = '=CEILING(B10,10)'; wp['B11'].fill = YELLOW; wp['B11'].font = BOLD
wp['E11'] = '90% повторных покупок происходят в пределах этого срока. Кто молчит дольше — «давно не покупал».'; wp['E11'].font = NOTE

wp['A13'] = 'F — частота: вероятность следующей покупки'; wp['A13'].font = BOLD
for j, h in enumerate(['Покупок (N)', 'Людей с ≥N', 'Сделали N+1', 'Доля следующей'], 1):
    c = wp.cell(row=14, column=j, value=h); c.font = H_FONT; c.fill = H_FILL
for k in range(1, 8):
    r = 14 + k
    wp.cell(row=r, column=1, value=k).font = BLUE
    wp.cell(row=r, column=2, value=f"=COUNTIF('Покупатели'!G:G,\">=\"&A{r})")
    wp.cell(row=r, column=3, value=f"=COUNTIF('Покупатели'!G:G,\">=\"&A{r}+1)")
    c = wp.cell(row=r, column=4, value=f'=IFERROR(C{r}/B{r},0)'); c.number_format = PCT
wp['A23'] = 'Порог F: «часто» — от N покупок'; wp['B23'] = 3; wp['B23'].font = Font(name=F, bold=True, color='0000FF'); wp['B23'].fill = YELLOW
wp['E23'] = 'Решение по таблице выше: после 3-й покупки доля следующей перестаёт расти (≈50%) — человек стал постоянным.'; wp['E23'].font = NOTE
wp['B23'].comment = Comment('Это решение аналитика, а не формула: выбираю N, после которого доля следующей покупки выходит на плато.', 'RFM')

wp['A25'] = 'M — деньги'; wp['A25'].font = BOLD
wp['A26'] = 'Медиана суммы за 12 месяцев'; wp['B26'] = "=MEDIAN('Покупатели'!J:J)"; wp['B26'].number_format = MONEY
wp['A27'] = '75-й перцентиль суммы'; wp['B27'] = "=PERCENTILE('Покупатели'!J:J,0.75)"; wp['B27'].number_format = MONEY
wp['A28'] = 'Порог M: верхние 25%, округлено до 100 ₸'; wp['B28'] = '=ROUND(B27,-2)'; wp['B28'].number_format = MONEY; wp['B28'].fill = YELLOW; wp['B28'].font = BOLD

# ---------------- Покупатели ----------------
wc = wb.create_sheet('Покупатели', 2)
header(wc, 1, ['Покупатель', 'Ниша', 'Канал', 'Первая покупка вообще', 'Строка с (в «Заказах»)', 'Строка по', 'F — заказов', 'Последняя покупка',
               'R — дней с последней', 'M — сумма за 12 мес.', 'Дней от 1-й до 2-й покупки', 'Когорта', 'Повтор за 30 дней (1/0)', 'Сегмент', 'Высокий M'],
       [12, 20, 14, 13, 12, 10, 10, 13, 11, 14, 13, 13, 12, 18, 10])
niches, chans = D['niches'], D['channels']
for i, (cid, ni, ch, fe) in enumerate(D['customers'], 2):
    a, b = rows_of[cid]
    wc.cell(row=i, column=1, value=cid).font = BLUE
    wc.cell(row=i, column=2, value=niches[ni]).font = BLUE
    wc.cell(row=i, column=3, value=chans[ch]).font = BLUE
    c = wc.cell(row=i, column=4, value=day2date(fe)); c.font = BLUE; c.number_format = DATE
    wc.cell(row=i, column=5, value=a).font = BLUE
    wc.cell(row=i, column=6, value=b).font = BLUE
    wc.cell(row=i, column=7, value=f'=F{i}-E{i}+1')
    c = wc.cell(row=i, column=8, value=f"=MAX(INDEX('Заказы'!$C:$C,E{i}):INDEX('Заказы'!$C:$C,F{i}))"); c.number_format = DATE
    wc.cell(row=i, column=9, value=f"='Пороги'!$B$3-H{i}")
    c = wc.cell(row=i, column=10, value=f"=SUM(INDEX('Заказы'!$D:$D,E{i}):INDEX('Заказы'!$D:$D,F{i}))"); c.number_format = MONEY
    wc.cell(row=i, column=11, value=f"=IF(G{i}>=2,INDEX('Заказы'!$C:$C,E{i}+1)-INDEX('Заказы'!$C:$C,E{i}),\"\")")
    c = wc.cell(row=i, column=12, value=f"=IF(D{i}>='Пороги'!$B$4,DATE(YEAR(D{i}),MONTH(D{i}),1),\"до окт 2025\")"); c.number_format = MONTH
    wc.cell(row=i, column=13, value=f"=IF(OR(L{i}=\"до окт 2025\",D{i}>'Пороги'!$B$3-30),\"\",IF(AND(G{i}>=2,K{i}<=30),1,0))")
    wc.cell(row=i, column=14, value=f"=IF(I{i}<='Пороги'!$B$11,IF(G{i}>='Пороги'!$B$23,\"Champions\",\"Potential Loyalists\"),IF(G{i}>='Пороги'!$B$23,\"At Risk\",\"Hibernating\"))")
    wc.cell(row=i, column=15, value=f"=IF(J{i}>='Пороги'!$B$28,\"да\",\"нет\")")
wc.freeze_panes = 'B2'
wc['E1'].comment = Comment('Заказы одного покупателя в листе «Заказы» идут подряд: с этой строки…', 'RFM')
wc['F1'].comment = Comment('…по эту строку. Формулы F, R и M берут именно этот диапазон.', 'RFM')

# ---------------- Сегменты ----------------
wsg = wb.create_sheet('Сегменты')
wsg['A1'] = 'Шаг 4. Сегменты на 30.09.2026'; wsg['A1'].font = TITLE
wsg['A2'] = "=\"Пороги: недавно — до \"&'Пороги'!B11&\" дней, часто — от \"&'Пороги'!B23&\" покупок, высокий M — от \"&TEXT('Пороги'!B28,\"# ##0\")&\" ₸\""
wsg['A2'].font = NOTE
cols = ['Сегмент', 'Правило', 'Покупателей', 'Доля людей', 'Оборот за 12 мес.', 'Доля оборота', 'Средний F', 'Средний R, дней', 'Средний M', 'С высоким M']
for j, h in enumerate(cols, 1):
    c = wsg.cell(row=4, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = Alignment(wrap_text=True)
for j, w in enumerate([22, 26, 13, 11, 18, 11, 10, 12, 13, 12], 1):
    wsg.column_dimensions[get_column_letter(j)].width = w
segs = [('Champions', 'недавно · часто'), ('Potential Loyalists', 'недавно · 1–2 покупки'), ('At Risk', 'давно · часто'), ('Hibernating', 'давно · 1–2 покупки')]
P = "'Покупатели'!"
for k, (s, rule) in enumerate(segs):
    r = 5 + k
    c = wsg.cell(row=r, column=1, value=s); c.font = BOLD; c.fill = PatternFill('solid', fgColor=SEGFILL[s])
    wsg.cell(row=r, column=2, value=rule)
    wsg.cell(row=r, column=3, value=f'=COUNTIF({P}N:N,A{r})').number_format = '#,##0'
    wsg.cell(row=r, column=4, value=f'=C{r}/$C$9').number_format = PCT
    wsg.cell(row=r, column=5, value=f'=SUMIF({P}N:N,A{r},{P}J:J)').number_format = MONEY
    wsg.cell(row=r, column=6, value=f'=E{r}/$E$9').number_format = PCT
    wsg.cell(row=r, column=7, value=f'=AVERAGEIF({P}N:N,A{r},{P}G:G)').number_format = '0.00'
    wsg.cell(row=r, column=8, value=f'=AVERAGEIF({P}N:N,A{r},{P}I:I)').number_format = '0'
    wsg.cell(row=r, column=9, value=f'=AVERAGEIF({P}N:N,A{r},{P}J:J)').number_format = MONEY
    wsg.cell(row=r, column=10, value=f'=COUNTIFS({P}N:N,A{r},{P}O:O,"да")').number_format = '#,##0'
wsg['A9'] = 'Итого'; wsg['A9'].font = BOLD
wsg['C9'] = '=SUM(C5:C8)'; wsg['C9'].number_format = '#,##0'; wsg['C9'].font = BOLD
wsg['D9'] = '=SUM(D5:D8)'; wsg['D9'].number_format = PCT
wsg['E9'] = '=SUM(E5:E8)'; wsg['E9'].number_format = MONEY; wsg['E9'].font = BOLD
wsg['F9'] = '=SUM(F5:F8)'; wsg['F9'].number_format = PCT
wsg['J9'] = '=SUM(J5:J8)'; wsg['J9'].number_format = '#,##0'
wsg['A11'] = 'Проверка: число покупателей в листе «Покупатели»'; wsg['C11'] = f'=COUNTA({P}A:A)-1'; wsg['C11'].number_format = '#,##0'
wsg['A12'] = 'Проверка: сумма заказов в листе «Заказы»'; wsg['E12'] = "=SUM('Заказы'!D:D)"; wsg['E12'].number_format = MONEY

# ---------------- По месяцам ----------------
wm = wb.create_sheet('По месяцам')
wm['A1'] = 'Шаг 5. Динамика по месяцам'; wm['A1'].font = TITLE
cols = ['Месяц', 'Заказы', 'Выручка', 'Новые покупатели', 'Покупатели в месяце', 'Повторные заказы', 'Доля повторных заказов', 'Средний чек', 'Повтор за 30 дней у новых']
for j, h in enumerate(cols, 1):
    c = wm.cell(row=3, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = Alignment(wrap_text=True)
for j, w in enumerate([12, 10, 16, 12, 12, 12, 12, 12, 14], 1):
    wm.column_dimensions[get_column_letter(j)].width = w
Z = "'Заказы'!"
for k in range(12):
    r = 4 + k
    y, m = (2025, 10 + k) if k < 3 else (2026, k - 2)
    c = wm.cell(row=r, column=1, value=datetime.date(y, m, 1)); c.number_format = MONTH; c.font = BLUE
    wm.cell(row=r, column=2, value=f'=COUNTIF({Z}F:F,A{r})').number_format = '#,##0'
    wm.cell(row=r, column=3, value=f'=SUMIF({Z}F:F,A{r},{Z}D:D)').number_format = MONEY
    wm.cell(row=r, column=4, value=f'=COUNTIFS({Z}F:F,A{r},{Z}E:E,1)').number_format = '#,##0'
    wm.cell(row=r, column=5, value=f'=COUNTIFS({Z}F:F,A{r},{Z}G:G,1)').number_format = '#,##0'
    wm.cell(row=r, column=6, value=f'=B{r}-D{r}').number_format = '#,##0'
    wm.cell(row=r, column=7, value=f'=F{r}/B{r}').number_format = PCT
    wm.cell(row=r, column=8, value=f'=C{r}/B{r}').number_format = MONEY
    c = wm.cell(row=r, column=9, value=f'=IFERROR(AVERAGEIF({P}L:L,A{r},{P}M:M),"ещё считается")'); c.number_format = PCT
wm['A17'] = 'Повтор за 30 дней: доля новых покупателей месяца, которые купили снова в течение 30 дней. Для сентября 2026 ещё не прошло 30 дней — поэтому пусто.'
wm['A17'].font = NOTE
wm['A18'] = '20.07.2026 в модели меняется шаблон воронки: после первой оплаты больше нет сообщений от Smart консультанта.'; wm['A18'].font = NOTE

# ---------------- Когорты ----------------
wk = wb.create_sheet('Когорты')
wk['A1'] = 'Шаг 6. Когорты: какая доля покупателей снова купила через N месяцев'; wk['A1'].font = TITLE
wk['A2'] = 'Месяц 0 — месяц первой покупки (всегда 100%). Месяц 1 — следующий календарный месяц, и т. д. Пусто — этот месяц ещё не наступил.'; wk['A2'].font = NOTE
hdr = ['Когорта', 'Размер'] + [f'Мес. {k}' for k in range(12)]
for j, h in enumerate(hdr, 1):
    c = wk.cell(row=4, column=j, value=h); c.font = H_FONT; c.fill = H_FILL
wk.column_dimensions['A'].width = 12; wk.column_dimensions['B'].width = 9
for j in range(3, 15):
    wk.column_dimensions[get_column_letter(j)].width = 8.5
for k in range(12):
    r = 5 + k
    y, m = (2025, 10 + k) if k < 3 else (2026, k - 2)
    c = wk.cell(row=r, column=1, value=datetime.date(y, m, 1)); c.number_format = MONTH; c.font = BLUE
    wk.cell(row=r, column=2, value=f'=COUNTIF({P}L:L,A{r})').number_format = '#,##0'
    for off in range(12):
        col = 3 + off
        if k + off <= 11:
            f = f'=COUNTIFS({Z}H:H,$A{r},{Z}F:F,EDATE($A{r},{off}),{Z}G:G,1)/$B{r}'
            c = wk.cell(row=r, column=col, value=f); c.number_format = '0%'
wk['A18'] = 'Как читать: строка «июл 2026», столбец «Мес. 1» — какая доля тех, кто впервые купил в июле, снова купил в августе.'; wk['A18'].font = NOTE
from openpyxl.formatting.rule import ColorScaleRule
wk.conditional_formatting.add('D5:N16', ColorScaleRule(start_type='num', start_value=0, start_color='FFFFFF', end_type='num', end_value=0.25, end_color='2A78D6'))

# ---------------- Пример ----------------
wx = wb.create_sheet('Пример')
wx['A1'] = 'Шаг 7. Пример вручную: 10 покупателей'; wx['A1'].font = TITLE
wx['A2'] = 'Слева — их заказы. Справа — R, F, M и сегмент простыми формулами по этим заказам.'; wx['A2'].font = NOTE
for j, h in enumerate(['Покупатель', 'Дата', 'Сумма'], 1):
    c = wx.cell(row=4, column=j, value=h); c.font = H_FONT; c.fill = H_FILL
# pick examples: deterministic variety
cust = D['customers']
order_by = {}
for oid, cid, d, amt in D['orders']:
    order_by.setdefault(cid, []).append((d, amt))
W1 = D['W1']
def seg_of(cid):
    os = order_by[cid]; R = W1 - os[-1][0]; Fq = len(os)
    return ('Champions' if Fq >= 3 else 'Potential Loyalists') if R <= 60 else ('At Risk' if Fq >= 3 else 'Hibernating')
picked = []
want = ['Champions', 'Champions', 'Potential Loyalists', 'Potential Loyalists', 'Potential Loyalists', 'At Risk', 'At Risk', 'Hibernating', 'Hibernating', 'Hibernating']
pool = [c[0] for c in cust[::97]]
for w in want:
    for cid in pool:
        if cid not in picked and seg_of(cid) == w and len(order_by[cid]) <= 6:
            picked.append(cid); break
r = 5
for cid in picked:
    for d, amt in order_by[cid]:
        wx.cell(row=r, column=1, value=cid).font = BLUE
        c = wx.cell(row=r, column=2, value=day2date(d)); c.number_format = DATE; c.font = BLUE
        c = wx.cell(row=r, column=3, value=amt); c.number_format = MONEY; c.font = BLUE
        r += 1
last = r - 1
for j, h in enumerate(['Покупатель', 'F = СЧЁТ заказов', 'Последняя покупка', 'R = дата анализа − последняя', 'M = СУММА', 'Сегмент'], 5):
    c = wx.cell(row=4, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = Alignment(wrap_text=True)
wx.row_dimensions[4].height = 32
for j, w in enumerate([12, 12, 12, 2, 12, 12, 14, 16, 14, 20], 1):
    wx.column_dimensions[get_column_letter(j)].width = w
for k, cid in enumerate(picked):
    rr = 5 + k
    wx.cell(row=rr, column=5, value=cid).font = BLUE
    wx.cell(row=rr, column=6, value=f'=COUNTIF($A$5:$A${last},E{rr})')
    c = wx.cell(row=rr, column=7, value=f'=_xlfn.MAXIFS($B$5:$B${last},$A$5:$A${last},E{rr})'); c.number_format = DATE
    wx.cell(row=rr, column=8, value=f"='Пороги'!$B$3-G{rr}")
    c = wx.cell(row=rr, column=9, value=f'=SUMIF($A$5:$A${last},E{rr},$C$5:$C${last})'); c.number_format = MONEY
    wx.cell(row=rr, column=10, value=f"=IF(H{rr}<='Пороги'!$B$11,IF(F{rr}>='Пороги'!$B$23,\"Champions\",\"Potential Loyalists\"),IF(F{rr}>='Пороги'!$B$23,\"At Risk\",\"Hibernating\"))")
wx['E17'] = 'Проверьте любого покупателя: найдите его в листе «Покупатели» — цифры совпадут.'; wx['E17'].font = NOTE

# ---------------- Стратегия ----------------
ST = json.load(open('/home/claude/portfolio/strat.json'))
wt = wb.create_sheet('Стратегия')
wt['A1'] = 'Шаг 8. Стратегия: откуда +12,7 млн ₸ выручки и −20% бюджета'; wt['A1'].font = TITLE
wt['A2'] = 'Сегмент каждого покупателя на 31.08.2026 и его покупки 1–30.09.2026 (посчитано по заказам, включая историю до окна анализа). Синие ячейки — данные и решения, чёрные — формулы.'; wt['A2'].font = NOTE
cols = ['Сегмент на 31.08', 'Людей', 'Купили в сентябре', 'Доля сейчас', 'Доля до 20 июля (июнь)', 'Цель', 'Почему цель', 'Доп. покупателей', 'Выручка на покупателя', 'Доп. выручка в месяц', 'Заказов на покупателя', 'Доп. заказов']
for j, h in enumerate(cols, 1):
    c = wt.cell(row=4, column=j, value=h); c.font = H_FONT; c.fill = H_FILL; c.alignment = Alignment(wrap_text=True, vertical='center')
wt.row_dimensions[4].height = 44
for j, w in enumerate([20, 9, 11, 10, 12, 9, 40, 12, 13, 16, 11, 11], 1):
    wt.column_dimensions[get_column_letter(j)].width = w
why = {'champ': 'вернуть уровень до 20 июля', 'loyal': 'вернуть уровень до 20 июля (сообщения после оплаты)', 'risk': 'гипотеза: персональные предложения, проверяю A/B', 'hib': 'не закладываю'}
names = {'champ': 'Champions', 'loyal': 'Potential Loyalists', 'risk': 'At Risk', 'hib': 'Hibernating'}
for i, k in enumerate(['loyal', 'risk', 'champ', 'hib']):
    r = 5 + i; d = ST[k]
    c = wt.cell(row=r, column=1, value=names[k]); c.font = BOLD; c.fill = PatternFill('solid', fgColor=SEGFILL[names[k]])
    wt.cell(row=r, column=2, value=d['n']).font = BLUE
    wt.cell(row=r, column=3, value=d['buy']).font = BLUE
    wt.cell(row=r, column=4, value=f'=C{r}/B{r}').number_format = PCT
    c = wt.cell(row=r, column=5, value=d['bbuy'] / d['bn']); c.font = BLUE; c.number_format = PCT
    if k == 'loyal': f = 0.18
    elif k == 'risk': f = 0.03
    elif k == 'champ': f = f'=E{r}'
    else: f = f'=D{r}'
    c = wt.cell(row=r, column=6, value=f); c.number_format = PCT; c.fill = YELLOW
    if isinstance(f, float): c.font = BLUE
    wt.cell(row=r, column=7, value=why[k])
    wt.cell(row=r, column=8, value=f'=B{r}*(F{r}-D{r})').number_format = '#,##0'
    c = wt.cell(row=r, column=9, value=d['rev'] / d['buy']); c.font = BLUE; c.number_format = MONEY
    wt.cell(row=r, column=10, value=f'=H{r}*I{r}').number_format = MONEY
    c = wt.cell(row=r, column=11, value=d['orders'] / d['buy']); c.font = BLUE; c.number_format = '0.00'
    wt.cell(row=r, column=12, value=f'=H{r}*K{r}').number_format = '#,##0'
wt['A9'] = 'Итого'; wt['A9'].font = BOLD
wt['H9'] = '=SUM(H5:H8)'; wt['H9'].number_format = '#,##0'; wt['H9'].font = BOLD
wt['J9'] = '=SUM(J5:J8)'; wt['J9'].number_format = MONEY; wt['J9'].font = BOLD
wt['L9'] = '=SUM(L5:L8)'; wt['L9'].number_format = '#,##0'
wt['F5'].comment = Comment('Июньский уровень Potential Loyalists — 18,2%. Беру 18%: возврат к уровню до изменения воронки.', 'RFM')

wt['A11'] = 'Выручка'; wt['A11'].font = BOLD
wt['A12'] = 'Выручка сентября 2026'; wt['C12'] = "='По месяцам'!C15"; wt['C12'].number_format = MONEY; wt['C12'].font = GREEN
wt['A13'] = 'Доп. выручка в месяц'; wt['C13'] = '=J9'; wt['C13'].number_format = MONEY
wt['A14'] = 'Выручка с целью'; wt['C14'] = '=C12+C13'; wt['C14'].number_format = MONEY; wt['C14'].font = BOLD
wt['A15'] = 'Рост выручки'; wt['C15'] = '=C13/C12'; wt['C15'].number_format = PCT; wt['C15'].font = BOLD

wt['A17'] = 'Маркетинговый бюджет в месяц'; wt['A17'].font = BOLD
for j, h in enumerate(['Направление', 'Было', 'Стало', 'Изменение', '', '', 'Почему'], 1):
    if h:
        c = wt.cell(row=18, column=j, value=h); c.font = H_FONT; c.fill = H_FILL
wt['A26'] = 'Бюджет всего (допущение ≈10% выручки)'; wt['B26'] = 11000000; wt['B26'].font = BLUE; wt['B26'].number_format = MONEY
wt['A27'] = 'Доля привлечения в «Было»'; wt['B27'] = 0.5; wt['B27'].font = BLUE; wt['B27'].number_format = '0%'
wt['A28'] = 'Остальное «Было» — удержание, делилось по числу людей (доли из листа «Сегменты» на 30.09)'; wt['A28'].font = NOTE
seg_row = {'Champions': 5, 'Potential Loyalists': 6, 'At Risk': 7, 'Hibernating': 8}
brows = [('Привлечение новых', '=$B$26*$B$27', 4400000, 'Сокращаю TikTok: повтор у его покупателей 16,5% против 22,7–29,8% в других каналах; часть денег — в look-alike по Champions'),
         ('Potential Loyalists', None, 2300000, 'главный рычаг: цепочка ко второй покупке, баллы'),
         ('At Risk', None, 1500000, 'персональные предложения; SMS — только высокому M'),
         ('Champions', None, 300000, 'скидки заменяю статусом и рекомендациями'),
         ('Hibernating', None, 300000, 'одно бесплатное сообщение; исключаю из платных аудиторий')]
for i, (nm, before, after, w) in enumerate(brows):
    r = 19 + i
    wt.cell(row=r, column=1, value=nm).font = BOLD
    f = before if before else f"=$B$26*(1-$B$27)*'Сегменты'!D{seg_row[nm]}"
    wt.cell(row=r, column=2, value=f).number_format = MONEY
    c = wt.cell(row=r, column=3, value=after); c.font = BLUE; c.number_format = MONEY
    wt.cell(row=r, column=4, value=f'=C{r}-B{r}').number_format = '+#,##0" ₸";-#,##0" ₸"'
    wt.cell(row=r, column=7, value=w)
wt['A24'] = 'Итого'; wt['A24'].font = BOLD
wt['B24'] = '=SUM(B19:B23)'; wt['B24'].number_format = MONEY; wt['B24'].font = BOLD
wt['C24'] = '=SUM(C19:C23)'; wt['C24'].number_format = MONEY; wt['C24'].font = BOLD
wt['D24'] = '=C24/B24-1'; wt['D24'].number_format = PCT; wt['D24'].font = BOLD

wt['A30'] = 'Стоимость заказа'; wt['A30'].font = BOLD
wt['A31'] = 'Заказов в сентябре'; wt['C31'] = "='По месяцам'!B15"; wt['C31'].font = GREEN; wt['C31'].number_format = '#,##0'
wt['A32'] = 'Маркетинг на заказ: было'; wt['C32'] = '=B24/C31'; wt['C32'].number_format = MONEY
wt['A33'] = 'Маркетинг на заказ: стало'; wt['C33'] = '=C24/(C31+L9)'; wt['C33'].number_format = MONEY
wt['A34'] = 'Изменение'; wt['C34'] = '=C33/C32-1'; wt['C34'].number_format = PCT; wt['C34'].font = BOLD

for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)


print([s.title for s in wb.worksheets])
wb.save('/home/claude/portfolio/RFM_SmartFlow_данные_12_месяцев.xlsx')
print('picked', picked)
